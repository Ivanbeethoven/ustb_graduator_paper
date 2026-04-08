from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


CANVAS_W = 16.8
CANVAS_H = 11.0
SLIDE_W = 13.333
SLIDE_H = 7.5
FONT_NAME = "Microsoft YaHei"

LAYER = {
    "L1": ("E3F2FD", "1565C0"),
    "L2": ("FFF3E0", "E65100"),
    "L3": ("F3E5F5", "6A1B9A"),
    "L4": ("E8F5E9", "2E7D32"),
}

COMP_FILL = "FFFFFF"
COMP_BORDER = "546E7A"
EXT_FILL = "ECEFF1"
EXT_BORDER = "78909C"
ARROW_CLR = "263238"
LBL_CLR = "37474F"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.replace("#", ""))


def sx(value: float):
    return Inches(value * SLIDE_W / CANVAS_W)


def sy(value: float):
    return Inches(value * SLIDE_H / CANVAS_H)


def px(value: float):
    return sx(value)


def py(value: float):
    return sy(CANVAS_H - value)


def rect(x0: float, y0: float, width: float, height: float):
    return px(x0), py(y0 + height), sx(width), sy(height)


def add_text(shape, blocks, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE):
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.vertical_anchor = valign
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0

    for idx, block in enumerate(blocks):
        paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        paragraph.alignment = block.get("align", align)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = block["text"]
        font = run.font
        font.name = block.get("font", FONT_NAME)
        font.size = Pt(block.get("size", 10))
        font.bold = block.get("bold", False)
        font.italic = block.get("italic", False)
        font.color.rgb = rgb(block.get("color", "1A1A2E"))


def add_textbox(slide, x0, y0, width, height, blocks, fill=None, line=None,
                radius=False, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE):
    if fill is None and line is None and not radius:
        shape = slide.shapes.add_textbox(*rect(x0, y0, width, height))
        shape.fill.background()
        shape.line.fill.background()
    else:
        auto_shape = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        shape = slide.shapes.add_shape(auto_shape, *rect(x0, y0, width, height))
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(0.8)
    add_text(shape, blocks, align=align, valign=valign)
    return shape


def ensure_tail_arrow(connector):
    line = connector.element.spPr.ln
    if line is None:
        return
    for child in list(line):
        if child.tag.endswith("tailEnd"):
            line.remove(child)
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "arrow")
    line.append(tail)


def add_segment(slide, start, end, color=ARROW_CLR, width=1.5, dashed=False, arrow=False):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        px(start[0]), py(start[1]),
        px(end[0]), py(end[1]),
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ensure_tail_arrow(connector)
    return connector


def add_label(slide, x, y, text, width=None, height=0.34, size=8.4):
    if width is None:
        width = max(1.15, min(2.1, 0.16 * len(text) + 0.35))
    add_textbox(
        slide,
        x - width / 2,
        y - height / 2,
        width,
        height,
        [{"text": text, "size": size, "color": LBL_CLR}],
        fill="FFFFFF",
        line="BDBDBD",
        radius=True,
    )


def add_polyline(slide, points, label=None, label_pos=None, color=ARROW_CLR, width=1.5, dashed=False, arrow=True):
    for idx in range(len(points) - 1):
        add_segment(
            slide,
            points[idx],
            points[idx + 1],
            color=color,
            width=width,
            dashed=dashed,
            arrow=arrow and idx == len(points) - 2,
        )
    if label and label_pos:
        add_label(slide, label_pos[0], label_pos[1], label)


def add_u_component_icon(slide, x0, y0, width, height, color=COMP_BORDER):
    bw, bh = 0.19, 0.23
    tw, th = 0.08, 0.066
    bx = x0 + width - bw - 0.08
    by = y0 + height - bh - 0.06

    main = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *rect(bx, by, bw, bh))
    main.fill.solid()
    main.fill.fore_color.rgb = rgb("FFFFFF")
    main.line.color.rgb = rgb(color)
    main.line.width = Pt(0.8)

    for ty in (by + bh - th - 0.02, by + th - 0.01):
        tab = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *rect(bx - tw / 2, ty, tw, th))
        tab.fill.solid()
        tab.fill.fore_color.rgb = rgb("FFFFFF")
        tab.line.color.rgb = rgb(color)
        tab.line.width = Pt(0.8)


def add_layer_band(slide, x0, y0, width, height, label_cn, label_en, fill_color, border_color):
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, *rect(x0, y0, width, height))
    band.fill.solid()
    band.fill.fore_color.rgb = rgb(fill_color)
    band.line.color.rgb = rgb(border_color)
    band.line.width = Pt(1.8)
    add_textbox(
        slide,
        x0 + 0.24,
        y0 + height - 0.44,
        4.6,
        0.34,
        [{"text": f"«layer»  {label_cn}  {label_en}", "size": 12.0, "bold": True, "color": border_color, "align": PP_ALIGN.LEFT}],
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
    )


def add_component(slide, cx, cy, width, height, name, sub="", border=COMP_BORDER, fill=COMP_FILL):
    x0 = cx - width / 2
    y0 = cy - height / 2
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *rect(x0, y0, width, height))
    card.fill.solid()
    card.fill.fore_color.rgb = rgb(fill)
    card.line.color.rgb = rgb(border)
    card.line.width = Pt(1.1)
    add_u_component_icon(slide, x0, y0, width, height, color=border)
    add_textbox(
        slide,
        x0 + 0.16,
        y0 + height - 0.24,
        width - 0.5,
        0.16,
        [{"text": "«component»", "size": 8.3, "italic": True, "color": "555555"}],
        valign=MSO_ANCHOR.TOP,
    )
    name_y = y0 + (0.36 if sub else 0.28)
    add_textbox(
        slide,
        x0 + 0.16,
        name_y,
        width - 0.32,
        0.24,
        [{"text": name, "size": 12.4, "bold": True, "color": "1A1A2E"}],
    )
    if sub:
        add_textbox(
            slide,
            x0 + 0.18,
            y0 + 0.12,
            width - 0.36,
            0.16,
            [{"text": sub, "size": 9.5, "color": "555555"}],
        )


def add_datastore(slide, cx, cy, width, height, name, sub="", fill="E8F5E9", border="2E7D32"):
    x0 = cx - width / 2
    y0 = cy - height / 2
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *rect(x0, y0, width, height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(border)
    shape.line.width = Pt(1.2)
    add_textbox(
        slide,
        x0 + 0.18,
        y0 + height - 0.28,
        width - 0.36,
        0.18,
        [{"text": "«artifact»", "size": 8.4, "italic": True, "color": border}],
        valign=MSO_ANCHOR.TOP,
    )
    add_textbox(
        slide,
        x0 + 0.18,
        y0 + 0.34,
        width - 0.36,
        0.24,
        [{"text": name, "size": 12.2, "bold": True, "color": "1A1A2E"}],
    )
    if sub:
        add_textbox(
            slide,
            x0 + 0.18,
            y0 + 0.13,
            width - 0.36,
            0.16,
            [{"text": sub, "size": 9.4, "color": "555555"}],
        )


def add_external(slide, cx, cy, width, height, name, sub=""):
    x0 = cx - width / 2
    y0 = cy - height / 2
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *rect(x0, y0, width, height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(EXT_FILL)
    shape.line.color.rgb = rgb(EXT_BORDER)
    shape.line.width = Pt(1.2)
    add_textbox(
        slide,
        x0 + 0.14,
        y0 + height - 0.22,
        width - 0.28,
        0.16,
        [{"text": "«node»", "size": 8.1, "italic": True, "color": EXT_BORDER}],
        valign=MSO_ANCHOR.TOP,
    )
    add_textbox(
        slide,
        x0 + 0.14,
        y0 + 0.34,
        width - 0.28,
        0.22,
        [{"text": name, "size": 11.6, "bold": True, "color": "1A1A2E"}],
    )
    if sub:
        add_textbox(
            slide,
            x0 + 0.14,
            y0 + 0.12,
            width - 0.28,
            0.16,
            [{"text": sub, "size": 9.0, "color": "555555"}],
        )


def add_bus_dot(slide, x, y):
    dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, *rect(x - 0.05, y - 0.05, 0.10, 0.10))
    dot.fill.solid()
    dot.fill.fore_color.rgb = rgb(ARROW_CLR)
    dot.line.color.rgb = rgb(ARROW_CLR)
    dot.line.width = Pt(0.6)


def build_slide():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── 层宽与起点 ──
    BX, BW = 0.3, 14.1

    # ── 四个层带 ──
    add_layer_band(slide, BX, 8.70, BW, 1.90, "展现层",    "Presentation Layer",  *LAYER["L1"])
    add_layer_band(slide, BX, 5.20, BW, 3.25, "服务层",    "Service Layer",       *LAYER["L2"])
    add_layer_band(slide, BX, 2.65, BW, 2.30, "智能引擎层", "Intelligence Layer",  *LAYER["L3"])
    add_layer_band(slide, BX, 0.40, BW, 2.00, "数据层",    "Data Layer",          *LAYER["L4"])

    # ────────────────────────────────
    # L1 展现层  (cy=9.65)  三等分
    # ────────────────────────────────
    L1_CY    = 9.65
    L1_H     = 0.88
    L1_W     = 3.0
    L1_X     = [2.6, 7.2, 11.8]          # 均匀分布

    add_component(slide, L1_X[0], L1_CY, L1_W, L1_H, "Vue3 视图组件")
    add_component(slide, L1_X[1], L1_CY, L1_W, L1_H, "Pinia 状态管理")
    add_component(slide, L1_X[2], L1_CY, L1_W, L1_H, "ECharts 可视化图表")

    # ────────────────────────────────
    # L2 服务层 上排：API网关 + 调度器    居中靠近
    # ────────────────────────────────
    API_CX,  API_CY  = 4.6, 7.75
    API_W,   API_H   = 3.8, 0.94
    SCH_CX,  SCH_CY  = 9.6, 7.75
    SCH_W,   SCH_H   = 3.2, 0.94

    add_component(slide, API_CX, API_CY, API_W, API_H, "REST API 网关", "Flask / 鉴权")
    add_component(slide, SCH_CX, SCH_CY, SCH_W, SCH_H, "任务调度器", "异步调度")

    # L2 下排：四个业务模块，等间距
    BIZ_CY = 6.20
    BIZ_W  = 2.7
    BIZ_H  = 0.90
    BIZ_X  = [2.0, 5.0, 8.0, 11.0]

    add_component(slide, BIZ_X[0], BIZ_CY, BIZ_W, BIZ_H, "CVE 知识库", "导入 / 检索")
    add_component(slide, BIZ_X[1], BIZ_CY, BIZ_W, BIZ_H, "策略生成", "候选生成")
    add_component(slide, BIZ_X[2], BIZ_CY, BIZ_W, BIZ_H, "评审与排序", "评审聚合")
    add_component(slide, BIZ_X[3], BIZ_CY, BIZ_W, BIZ_H, "部署管理", "下发 / 回滚")

    # ────────────────────────────────
    # L3 智能引擎层  间距均匀化
    # ────────────────────────────────
    KPE_CX, KPE_CY, KPE_W, KPE_H = 2.8, 3.80, 3.8, 1.08
    RAG_CX, RAG_CY, RAG_W, RAG_H = 6.9, 3.80, 2.5, 0.90
    COT_CX, COT_CY, COT_W, COT_H = 9.7, 3.80, 2.5, 0.90
    FBK_CX, FBK_CY, FBK_W, FBK_H = 12.5, 3.80, 2.5, 0.90

    add_component(slide, KPE_CX, KPE_CY, KPE_W, KPE_H, "KPEAgent 核心引擎", "统一编排",
                  border="6A1B9A", fill="F9F0FF")
    add_component(slide, RAG_CX, RAG_CY, RAG_W, RAG_H, "RAG 检索器", "证据检索")
    add_component(slide, COT_CX, COT_CY, COT_W, COT_H, "CoT 推理机", "漏洞分析")
    add_component(slide, FBK_CX, FBK_CY, FBK_W, FBK_H, "反馈优化器", "反馈修正")

    # ────────────────────────────────
    # L4 数据层  向中心靠拢
    # ────────────────────────────────
    DS1_CX, DS1_CY = 4.8,  1.35
    DS2_CX, DS2_CY = 9.8,  1.35
    DS_W, DS_H     = 3.2,  1.12

    add_datastore(slide, DS1_CX, DS1_CY, DS_W, DS_H, "文件存储", "CVE / 策略结果")
    add_datastore(slide, DS2_CX, DS2_CY, DS_W, DS_H, "向量索引", "语义检索")

    # ── 外部系统 ──
    EXT_CX, EXT_CY = 15.0, 6.20
    add_external(slide, EXT_CX, EXT_CY, 2.35, 1.12, "目标集群", "/ 主机")

    # ════════════════════════════════
    # 接 口 箭 头
    # ════════════════════════════════

    # A1: Vue3 → REST API  "HTTP / JSON"
    add_polyline(slide,
                 [(L1_X[0], L1_CY - L1_H / 2),
                  (L1_X[0], 8.55),
                  (API_CX,  8.55),
                  (API_CX,  API_CY + API_H / 2)],
                 label="HTTP / JSON", label_pos=(3.6, 8.55))

    # A2: REST API → 调度器
    add_polyline(slide,
                 [(API_CX + API_W / 2, API_CY),
                  (SCH_CX - SCH_W / 2, SCH_CY)])

    # A3: 调度器 → 策略生成  "任务分发"
    add_polyline(slide,
                 [(SCH_CX, SCH_CY - SCH_H / 2),
                  (SCH_CX, 7.05),
                  (BIZ_X[1], 7.05),
                  (BIZ_X[1], BIZ_CY + BIZ_H / 2)],
                 label="任务分发", label_pos=(7.3, 7.05))

    # A4: 策略生成 → KPEAgent  "生成调用"
    add_polyline(slide,
                 [(BIZ_X[1], BIZ_CY - BIZ_H / 2),
                  (BIZ_X[1], 5.08),
                  (KPE_CX,  5.08),
                  (KPE_CX,  KPE_CY + KPE_H / 2)],
                 label="生成调用", label_pos=(3.9, 5.08))

    # A5: 反馈优化器 → 文件存储  "持久化"
    add_polyline(slide,
                 [(FBK_CX,  FBK_CY - FBK_H / 2),
                  (FBK_CX,  2.70),
                  (DS1_CX,  2.70),
                  (DS1_CX,  DS1_CY + DS_H / 2)],
                 label="持久化", label_pos=(8.7, 2.70))

    # A6: RAG → 向量索引  "语义检索"  虚线
    add_polyline(slide,
                 [(RAG_CX + RAG_W / 2, RAG_CY - 0.10),
                  (RAG_CX + RAG_W / 2 + 0.40, RAG_CY - 0.10),
                  (RAG_CX + RAG_W / 2 + 0.40, 2.48),
                  (DS2_CX,   2.48),
                  (DS2_CX,   DS2_CY + DS_H / 2)],
                 label="语义检索", label_pos=(9.6, 2.48), dashed=True)

    # A7: 部署管理 → 目标集群  "K8s API / SSH"
    add_polyline(slide,
                 [(BIZ_X[3] + BIZ_W / 2, BIZ_CY),
                  (EXT_CX - 2.35 / 2,    EXT_CY)],
                 label="K8s API / SSH", label_pos=(13.2, 6.43))

    return prs


def main():
    output = Path(__file__).with_name("system_architecture.pptx")
    presentation = build_slide()
    presentation.save(output)
    print(f"saved {output.name}")


if __name__ == "__main__":
    main()