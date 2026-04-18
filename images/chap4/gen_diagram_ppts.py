from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = 13.333
SLIDE_H = 7.5
FONT_NAME = "Microsoft YaHei"
SMALL_FOUR_PT = 12  # 小四号

BASE_DIR = Path(__file__).resolve().parent

DIAGRAMS: list[dict[str, str]] = [
    {"stem": "cve_lifecycle", "title": "CVE 漏洞处置生命周期图"},
    {"stem": "data_flow_audit", "title": "数据流与审计归档图"},
    {"stem": "role_use_case", "title": "角色与用例关系图"},
    {"stem": "storage_schema", "title": "核心实体与存储结构图"},
    {"stem": "strategy_code_flow", "title": "策略生成与反馈优化流程图"},
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.replace("#", ""))


def add_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.68), Inches(0.28), Inches(12.0), Inches(0.45))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    font = run.font
    font.name = FONT_NAME
    font.size = Pt(SMALL_FOUR_PT)
    font.bold = False
    font.color.rgb = rgb("1F2937")


def add_frame(slide, left: float, top: float, width: float, height: float) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb("FFFFFF")
    shape.line.color.rgb = rgb("D1D5DB")
    shape.line.width = Pt(1.0)
    # Keep corners visually subtle.
    try:
        shape.adjustments[0] = 0.02
    except Exception:
        pass


def contain_size(img_path: Path, box_w: float, box_h: float) -> tuple[float, float]:
    with Image.open(img_path) as img:
        w, h = img.size
    aspect = w / h
    box_aspect = box_w / box_h
    if aspect >= box_aspect:
        width = box_w
        height = width / aspect
    else:
        height = box_h
        width = height * aspect
    return width, height


def create_ppt_for_diagram(stem: str, title: str) -> Path:
    png_path = BASE_DIR / f"{stem}.png"
    if not png_path.exists():
        raise FileNotFoundError(f"Missing PNG: {png_path}")

    pptx_path = BASE_DIR / f"{stem}.pptx"

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # White page background for thesis-style consistency.
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")

    add_title(slide, title)

    frame_left = 0.55
    frame_top = 0.82
    frame_width = 12.23
    frame_height = 6.05
    add_frame(slide, frame_left, frame_top, frame_width, frame_height)

    image_margin = 0.18
    box_w = frame_width - image_margin * 2
    box_h = frame_height - image_margin * 2
    img_w, img_h = contain_size(png_path, box_w, box_h)
    img_left = frame_left + (frame_width - img_w) / 2
    img_top = frame_top + (frame_height - img_h) / 2

    slide.shapes.add_picture(
        str(png_path),
        Inches(img_left),
        Inches(img_top),
        width=Inches(img_w),
        height=Inches(img_h),
    )

    prs.save(pptx_path)
    return pptx_path


def generate_all(diagrams: Iterable[dict[str, str]]) -> list[Path]:
    outputs: list[Path] = []
    for item in diagrams:
        outputs.append(create_ppt_for_diagram(item["stem"], item["title"]))
    return outputs


def main() -> None:
    outputs = generate_all(DIAGRAMS)
    for path in outputs:
        print(f"saved {path.name}")


if __name__ == "__main__":
    main()
