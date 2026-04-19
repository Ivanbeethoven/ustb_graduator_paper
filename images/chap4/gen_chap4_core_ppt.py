from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SLIDE_W = 13.333
SLIDE_H = 7.5
FONT_NAME = "Microsoft YaHei"
FONT_SIZE = 13
SMALL = 11.5
TINY = 9.5

INK = "1E1E1E"
LINE = "4A4A4A"
MID = "7A7A7A"
LIGHT = "B8B8B8"
SOFT = "F6F6F6"
WHITE = "FFFFFF"

BASE = Path(__file__).resolve().parent
OUT_COMBINED = BASE / "chap4_core_diagrams.pptx"
OUT_FLOW = BASE / "chap4_flowchart.pptx"
OUT_DATA = BASE / "chap4_dataflow.pptx"
OUT_USECASE = BASE / "chap4_usecase.pptx"
OUT_SCHEMA = BASE / "storage_schema.pptx"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.replace("#", ""))


def add_text(shape, blocks: list[dict], *, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, block in enumerate(blocks):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = block.get("align", align)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = block["text"]
        f = r.font
        f.name = FONT_NAME
        f.size = Pt(block.get("size", FONT_SIZE))
        f.bold = block.get("bold", False)
        f.italic = block.get("italic", False)
        f.color.rgb = rgb(block.get("color", INK))


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str | list[dict],
    *,
    shape_type=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    fill=WHITE,
    line=LINE,
    line_width=1.1,
    font_size=FONT_SIZE,
    bold=False,
    align=PP_ALIGN.CENTER,
    valign=MSO_ANCHOR.MIDDLE,
):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    runs = text if isinstance(text, list) else [{"text": text, "size": font_size, "bold": bold}]
    add_text(shape, runs, align=align, valign=valign)
    return shape


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str | list[dict],
    *,
    font_size=FONT_SIZE,
    align=PP_ALIGN.CENTER,
    valign=MSO_ANCHOR.MIDDLE,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    runs = text if isinstance(text, list) else [{"text": text, "size": font_size}]
    add_text(shape, runs, align=align, valign=valign)
    return shape


def ensure_arrow(connector) -> None:
    line = connector.element.spPr.ln
    if line is None:
        return
    for child in list(line):
        if child.tag.endswith("tailEnd"):
            line.remove(child)
    node = OxmlElement("a:tailEnd")
    node.set("type", "arrow")
    line.append(node)


def add_segment(slide, start: tuple[float, float], end: tuple[float, float], *, color=LINE, width=1.0, dashed=False, arrow=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start[0]),
        Inches(start[1]),
        Inches(end[0]),
        Inches(end[1]),
    )
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    if dashed:
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ensure_arrow(conn)
    return conn


def add_polyline(
    slide,
    points: list[tuple[float, float]],
    *,
    color=LINE,
    width=1.0,
    dashed=False,
    arrow=True,
    label: str | None = None,
    label_pos: tuple[float, float] | None = None,
    label_w: float = 1.2,
):
    for idx in range(len(points) - 1):
        add_segment(
            slide,
            points[idx],
            points[idx + 1],
            color=color,
            width=width,
            dashed=dashed,
            arrow=arrow and idx == len(points) - 2,
        )
    if label and label_pos:
        add_box(
            slide,
            label_pos[0] - label_w / 2,
            label_pos[1] - 0.14,
            label_w,
            0.28,
            label,
            shape_type=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            fill=WHITE,
            line=LIGHT,
            line_width=0.8,
            font_size=TINY,
        )


def add_actor(slide, cx: float, cy: float, label: str):
    head = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(cx - 0.13), Inches(cy - 0.78), Inches(0.26), Inches(0.26))
    head.fill.background()
    head.line.color.rgb = rgb(LINE)
    head.line.width = Pt(1.0)
    add_segment(slide, (cx, cy - 0.50), (cx, cy + 0.05), color=LINE, width=1.0)
    add_segment(slide, (cx - 0.20, cy - 0.25), (cx + 0.20, cy - 0.25), color=LINE, width=1.0)
    add_segment(slide, (cx, cy + 0.05), (cx - 0.18, cy + 0.38), color=LINE, width=1.0)
    add_segment(slide, (cx, cy + 0.05), (cx + 0.18, cy + 0.38), color=LINE, width=1.0)
    add_textbox(slide, cx - 0.8, cy + 0.42, 1.6, 0.32, label, font_size=SMALL)


def add_cylinder(slide, x: float, y: float, w: float, h: float, title: str):
    body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y + 0.10), Inches(w), Inches(h - 0.20))
    body.fill.solid()
    body.fill.fore_color.rgb = rgb(SOFT)
    body.line.color.rgb = rgb(LINE)
    body.line.width = Pt(1.0)
    top = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(0.24))
    top.fill.solid()
    top.fill.fore_color.rgb = rgb(SOFT)
    top.line.color.rgb = rgb(LINE)
    top.line.width = Pt(1.0)
    bottom = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y + h - 0.24), Inches(w), Inches(0.24))
    bottom.fill.solid()
    bottom.fill.fore_color.rgb = rgb(SOFT)
    bottom.line.color.rgb = rgb(LINE)
    bottom.line.width = Pt(1.0)
    add_textbox(slide, x + 0.08, y + 0.16, w - 0.16, h - 0.22, title, font_size=SMALL)


def add_class_box(slide, x: float, y: float, w: float, h: float, title: str, attrs: list[str], methods: list[str]):
    outer = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    outer.fill.solid()
    outer.fill.fore_color.rgb = rgb(WHITE)
    outer.line.color.rgb = rgb(MID)
    outer.line.width = Pt(1.0)

    header_h = 0.46
    attr_h = max(0.95, h * 0.48)
    method_y = y + header_h + attr_h
    add_segment(slide, (x, y + header_h), (x + w, y + header_h), color=LIGHT, width=0.9)
    add_segment(slide, (x, method_y), (x + w, method_y), color=LIGHT, width=0.9)

    add_textbox(slide, x + 0.06, y + 0.03, w - 0.12, header_h - 0.04, title, font_size=FONT_SIZE, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        x + 0.10,
        y + header_h + 0.07,
        w - 0.20,
        attr_h - 0.10,
        [{"text": "\n".join(attrs), "size": SMALL, "align": PP_ALIGN.LEFT}],
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
    )
    add_textbox(
        slide,
        x + 0.10,
        method_y + 0.07,
        w - 0.20,
        h - (method_y - y) - 0.10,
        [{"text": "\n".join(methods), "size": SMALL, "align": PP_ALIGN.LEFT}],
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
    )


def add_slide_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)


def build_flowchart(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)

    # group frame
    add_box(slide, 2.55, 0.22, 8.25, 6.90, "", shape_type=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, fill=WHITE, line=LIGHT, line_width=0.8)

    add_box(slide, 5.35, 0.36, 2.60, 0.46, "选择漏洞并发起分析会话", fill=WHITE, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 5.45, 1.16, 2.40, 0.42, "构建安全上下文", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 5.30, 1.92, 2.70, 0.50, "生成候选监控 / 防护策略", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 5.18, 2.88, 2.95, 0.58, "专家三维评审与同类型排序", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 5.55, 3.92, 2.20, 0.92, "是否需要\n反馈修订", fill=WHITE, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.DIAMOND)
    add_box(slide, 2.95, 5.12, 2.50, 0.52, "反馈优化并重新生成", fill=SOFT, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 5.50, 5.18, 2.30, 0.50, "确认交付结果", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 5.35, 6.05, 2.60, 0.50, "部署验证并记录审计", fill=WHITE, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)

    add_polyline(slide, [(6.65, 0.82), (6.65, 1.16)], arrow=True)
    add_polyline(slide, [(6.65, 1.58), (6.65, 1.92)], arrow=True)
    add_polyline(slide, [(6.65, 2.42), (6.65, 2.88)], arrow=True)
    add_polyline(slide, [(6.65, 3.46), (6.65, 3.92)], arrow=True)
    add_polyline(slide, [(6.65, 4.84), (6.65, 5.18)], label="否", label_pos=(6.25, 5.02), label_w=0.55, arrow=True)
    add_polyline(slide, [(6.65, 5.68), (6.65, 6.05)], arrow=True)
    add_polyline(slide, [(5.55, 4.38), (4.18, 4.38), (4.18, 5.12)], label="是", label_pos=(4.48, 4.08), label_w=0.55, arrow=True)
    add_polyline(slide, [(5.45, 5.38), (5.00, 5.38), (5.00, 2.18), (5.30, 2.18)], label="修订后返回生成阶段", label_pos=(4.38, 2.62), label_w=1.55, dashed=True, color=MID, arrow=True)


def build_dataflow(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)

    # Actors
    add_box(slide, 0.38, 0.40, 2.0, 0.42, "安全工程师", fill=WHITE, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 10.95, 0.40, 2.0, 0.42, "专家评审员", fill=WHITE, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)

    # Core process lane
    add_box(slide, 3.10, 0.26, 7.10, 6.72, "", fill=WHITE, line=LIGHT, line_width=0.8, shape_type=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)

    add_box(slide, 4.10, 0.72, 2.05, 0.40, "漏洞知识与证据管理", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 4.10, 1.70, 2.05, 0.46, "安全上下文构建", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 4.10, 2.76, 2.05, 0.54, "结构化漏洞分析与\n候选策略生成", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 7.00, 2.76, 2.00, 0.54, "三维评审与\n同类型排序", fill=WHITE, line=MID, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 7.00, 4.08, 2.00, 0.46, "反馈优化", fill=SOFT, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    add_box(slide, 7.00, 5.22, 2.00, 0.46, "确认结果与部署验证", fill=WHITE, line=LINE, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE)

    # Data stores
    add_cylinder(slide, 0.55, 1.48, 1.65, 0.66, "CVE 与多源证据")
    add_cylinder(slide, 0.55, 3.05, 1.65, 0.66, "候选策略与\n中间结果")
    add_cylinder(slide, 0.55, 5.15, 1.65, 0.66, "排序与审计记录")

    add_polyline(slide, [(2.38, 0.61), (4.10, 0.92)], label="发起会话 / 查询", label_pos=(3.10, 0.56), label_w=1.1, arrow=True)
    add_polyline(slide, [(6.15, 3.03), (7.00, 3.03)], label="候选提交评审", label_pos=(6.62, 2.62), label_w=1.0, arrow=True)
    add_polyline(slide, [(10.95, 0.61), (9.00, 3.03)], label="评审意见", label_pos=(9.92, 1.72), label_w=0.92, arrow=True)
    add_polyline(slide, [(9.00, 4.31), (7.85, 4.31), (7.85, 3.30), (6.15, 3.30)], label="修订建议", label_pos=(7.35, 3.72), label_w=0.92, arrow=True)
    add_polyline(slide, [(9.00, 5.45), (2.20, 5.48)], label="写入审计与排序记录", label_pos=(5.50, 5.18), label_w=1.42, arrow=True)
    add_polyline(slide, [(2.20, 1.82), (4.10, 1.90)], label="检索证据", label_pos=(3.05, 1.52), label_w=0.85, dashed=True, color=MID, arrow=True)
    add_polyline(slide, [(2.20, 3.38), (4.10, 3.10)], label="读取候选与中间结果", label_pos=(3.25, 3.38), label_w=1.35, dashed=True, color=MID, arrow=True)


def build_usecase(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)

    add_actor(slide, 1.02, 2.18, "安全工程师")
    add_actor(slide, 1.02, 5.28, "专家评审员")

    system = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.25), Inches(0.45), Inches(8.85), Inches(6.45))
    system.fill.solid()
    system.fill.fore_color.rgb = rgb(WHITE)
    system.line.color.rgb = rgb(LIGHT)
    system.line.width = Pt(0.8)

    usecases = {
        "浏览漏洞知识与证据": (3.15, 1.18, 1.60, 0.42),
        "创建分析会话": (3.15, 2.02, 1.30, 0.42),
        "生成候选策略": (5.55, 2.02, 1.30, 0.42),
        "反馈修订与再生成": (3.15, 2.92, 1.60, 0.42),
        "三维质量评审": (3.15, 4.38, 1.42, 0.42),
        "同类型候选排序": (3.15, 5.22, 1.48, 0.42),
        "确认交付结果": (5.55, 5.22, 1.42, 0.42),
        "部署验证": (8.10, 5.22, 1.15, 0.42),
        "查看评审与审计记录": (5.40, 6.05, 1.75, 0.42),
    }
    for label, (x, y, w, h) in usecases.items():
        add_box(slide, x, y, w, h, label, shape_type=MSO_AUTO_SHAPE_TYPE.OVAL, fill=WHITE, line=MID, font_size=SMALL)

    for s, e in [
        ((1.42, 1.82), (3.15, 1.39)),
        ((1.42, 2.00), (3.15, 2.23)),
        ((1.42, 2.25), (3.15, 3.13)),
        ((1.42, 2.45), (5.55, 5.43)),
        ((1.42, 4.98), (3.15, 4.59)),
        ((1.42, 5.18), (3.15, 5.43)),
        ((1.42, 5.42), (5.40, 6.26)),
    ]:
        add_segment(slide, s, e, color=LINE, width=0.9)

    add_polyline(slide, [(4.45, 2.23), (5.55, 2.23)], label="触发", label_pos=(5.00, 1.88), label_w=0.72, dashed=True, color=MID, arrow=True)
    add_polyline(slide, [(4.75, 3.13), (5.55, 2.44)], label="反馈驱动", label_pos=(5.05, 2.72), label_w=0.88, dashed=True, color=MID, arrow=True)
    add_polyline(slide, [(6.97, 5.43), (8.10, 5.43)], label="通过后执行", label_pos=(7.55, 5.08), label_w=1.02, dashed=True, color=MID, arrow=True)


def build_storage_schema(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)

    add_class_box(
        slide, 0.18, 1.50, 2.48, 2.60, "CVE",
        [
            "- String cve_id",
            "- String summary",
            "- String description",
            "- Float cvss_score",
            "- String severity",
            "- Date published_date",
            "- Enum status",
        ],
        [
            "+ getSeverityLevel(): String",
            "+ toFeatureVector(): List",
        ],
    )
    add_class_box(
        slide, 3.55, 1.72, 3.02, 2.38, "Session",
        [
            "- String session_id",
            "- String model_name",
            "- String prompt_template",
            "- DateTime created_at",
            "- Enum status",
        ],
        [
            "+ initWorkspace(): void",
            "+ getLatestStrategy(): Strategy",
        ],
    )
    add_class_box(
        slide, 7.22, 2.05, 2.38, 2.06, "Strategy",
        [
            "- String strategy_id",
            "- String content",
            "- String reasoning_chain",
            "- List<String> rag_snippets",
            "- Int version",
            "- DateTime generated_at",
        ],
        [
            "+ render(): String",
            "+ diff(other): String",
        ],
    )
    add_class_box(
        slide, 10.53, 1.56, 2.50, 2.50, "Ranking",
        [
            "- String ranking_id",
            "- String user_id",
            "- DateTime created_at",
            "- List<String> order",
            "- Map<String, Float> dimension_scores",
            "- String comment",
        ],
        [
            "+ validate(): Boolean",
            "+ toPreferencePair(): Pair",
        ],
    )

    add_polyline(slide, [(2.66, 2.80), (3.55, 2.80)], arrow=True)
    add_textbox(slide, 2.68, 2.48, 0.80, 0.20, "发起分析", font_size=SMALL)
    add_textbox(slide, 2.72, 2.95, 0.18, 0.16, "1", font_size=TINY)
    add_textbox(slide, 3.20, 2.48, 0.32, 0.16, "0..*", font_size=TINY)

    add_polyline(slide, [(6.57, 3.08), (7.22, 3.08)], arrow=True)
    add_textbox(slide, 6.58, 2.76, 0.55, 0.20, "生成", font_size=SMALL)
    add_textbox(slide, 6.54, 3.24, 0.18, 0.16, "1", font_size=TINY)
    add_textbox(slide, 6.93, 2.76, 0.36, 0.16, "1..*", font_size=TINY)

    add_polyline(slide, [(6.57, 2.18), (6.57, 1.22), (10.53, 1.22), (10.53, 2.04)], arrow=True)
    add_textbox(slide, 8.10, 0.94, 0.82, 0.20, "收集评价", font_size=SMALL)
    add_textbox(slide, 6.62, 1.92, 0.18, 0.16, "1", font_size=TINY)
    add_textbox(slide, 10.22, 0.94, 0.32, 0.16, "0..*", font_size=TINY)

    add_polyline(slide, [(10.53, 3.58), (9.60, 3.58)], arrow=True)
    add_textbox(slide, 9.80, 3.26, 0.82, 0.20, "排序对象", font_size=SMALL)
    add_textbox(slide, 9.42, 3.74, 0.34, 0.16, "1..*", font_size=TINY)
    add_textbox(slide, 10.03, 3.74, 0.30, 0.16, "0..*", font_size=TINY)


def build_presentation(builder):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    builder(prs)
    return prs


def main() -> None:
    flow = build_presentation(build_flowchart)
    flow.save(OUT_FLOW)

    data = build_presentation(build_dataflow)
    data.save(OUT_DATA)

    usecase = build_presentation(build_usecase)
    usecase.save(OUT_USECASE)

    schema = build_presentation(build_storage_schema)
    schema.save(OUT_SCHEMA)

    combined = Presentation()
    combined.slide_width = Inches(SLIDE_W)
    combined.slide_height = Inches(SLIDE_H)
    build_flowchart(combined)
    build_dataflow(combined)
    build_usecase(combined)
    build_storage_schema(combined)
    combined.save(OUT_COMBINED)
    print(f"saved {OUT_COMBINED.name}")


if __name__ == "__main__":
    main()
